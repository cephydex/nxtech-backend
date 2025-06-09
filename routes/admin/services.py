import logging
import openpyxl, io
from repos.member import MemberRepo

logger = logging.getLogger(__name__)


def process_new_member_id(old_member_id:str, prefix:str):
    num_part = int(old_member_id[2:])
    logger.warning("PROC 1: mID %s -- %d" % (old_member_id, num_part))
    num_part += 1
    num_str = f'{num_part:05d}'
    new_id = prefix + num_str
    logger.warning("PROC 2: mID %s -- %d" % (new_id, num_part))
    return new_id


def getNextMemberId():
    last_member = MemberRepo.get_last_member_id()
    # print("Members ORD", last_member)
    member_id = None
    prefix = "MG"
    if last_member:
        old_id = last_member['member_id']
        member_id = process_new_member_id(old_id, prefix)
        
    else:
        member_id = f"{prefix}00001"
    
    return member_id


async def prep_xlsx_file(file, sheet_name: str = None):
    f = await file.read()
    xlsx = io.BytesIO(f)
    
    # wb = openpyxl.load_workbook(xlsx)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    if sheet_name != None:
        ws = wb['Sheet1']
        
    # print("OPEN1", wb)
    df = list(ws.iter_rows(values_only=True))
    
    return df


def processAllName(all_names):
    name_parts = all_names.split(" ")
    first_name = ''; other_names = None
    if len(name_parts) == 1:
        first_name = name_parts[-1]
        del name_parts[-1]
    # print("All names", all_names, name_parts)
    # first_name = ''; other_names = None
    if len(name_parts):
        first_name = name_parts[0]
        del name_parts[0]
        other_names = ''
        if len(name_parts) > 0:
            other_names = " ".join(name_parts)
        
    
    # print(f"PROC :: FN: {first_name}, ON: {other_names}")
    return first_name, other_names

def processAllName2(all_names):
    name_parts = all_names.split(" ")
    del name_parts[-1]
    print("All names", all_names, name_parts)
    first_name = name_parts[0]
    
    other_temp = ''
    print('ARR LEN', len(name_parts))
    if len(name_parts) > 2:
        other_temp = name_parts[-1]
    other_names = other_temp if other_temp != '' else None






# print(f'{n:05d}')
# print('%05d' % n)
# print('{0:05d}'.format(n))